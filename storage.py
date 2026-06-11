"""
Save results to CSV, XLSX, and failed_urls.csv.
"""

import csv
import logging
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

import config

logger = logging.getLogger(__name__)


def save_results(records: list[dict], failed: list[dict]) -> None:
    """Write products_output.xlsx, products_output.csv, failed_urls.csv."""

    _save_products_csv(records)
    _save_products_xlsx(records)
    _save_failed_csv(failed)

    logger.info("Saved %d products → %s / %s", len(records),
                config.OUTPUT_CSV, config.OUTPUT_XLSX)
    logger.info("Saved %d failed URLs → %s", len(failed), config.FAILED_CSV)


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def _save_products_csv(records: list[dict]) -> None:
    path = Path(config.OUTPUT_CSV)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=config.OUTPUT_COLUMNS,
                                extrasaction="ignore")
        writer.writeheader()
        writer.writerows(records)


# ---------------------------------------------------------------------------
# XLSX with basic formatting
# ---------------------------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_ALT_FILL    = PatternFill("solid", fgColor="D6E4F0")


def _save_products_xlsx(records: list[dict]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"

    # Header row
    ws.append(config.OUTPUT_COLUMNS)
    for cell in ws[1]:
        cell.fill   = _HEADER_FILL
        cell.font   = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    for i, rec in enumerate(records, start=2):
        row = [rec.get(col, "") for col in config.OUTPUT_COLUMNS]
        ws.append(row)
        if i % 2 == 0:
            for cell in ws[i]:
                cell.fill = _ALT_FILL

    # Auto-width
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    ws.freeze_panes = "A2"
    wb.save(config.OUTPUT_XLSX)


# ---------------------------------------------------------------------------
# Failed URLs
# ---------------------------------------------------------------------------

def _save_failed_csv(failed: list[dict]) -> None:
    path = Path(config.FAILED_CSV)
    fieldnames = ["Brand Name", "URL", "Reason"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(failed)


# ---------------------------------------------------------------------------
# Incremental append (called after each product so progress is never lost)
# ---------------------------------------------------------------------------

_csv_initialized = False


def append_record(record: dict) -> None:
    global _csv_initialized
    path = Path(config.OUTPUT_CSV)
    write_header = not _csv_initialized or not path.exists()
    with path.open("a", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=config.OUTPUT_COLUMNS,
                                extrasaction="ignore")
        if write_header:
            writer.writeheader()
        writer.writerow(record)
    _csv_initialized = True
